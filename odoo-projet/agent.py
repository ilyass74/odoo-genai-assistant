import os
from pathlib import Path
from dotenv import load_dotenv
from docx import Document
from langchain_groq import ChatGroq
from langchain_core.tools import tool
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from connexion import get_stock, get_sales, get_invoices

load_dotenv("secrets.env")

embeddings = HuggingFaceEmbeddings(model_name="BAAI/bge-small-en-v1.5")
vectorstore = Chroma(
    collection_name="business_documents",
    embedding_function=embeddings,
    persist_directory="chroma_db"
)


@tool
def stock():
    """Retourne le stock actuel des produits."""
    return get_stock(limit=20)


@tool
def sales():
    """Retourne les dernieres commandes de vente."""
    return get_sales(limit=20)


@tool
def invoices():
    """Retourne les dernieres factures."""
    return get_invoices(limit=20)


@tool
def chercher_documents(question: str):
    """Recherche une information dans les documents internes de l''entreprise (rapports, contrats, politiques)."""
    resultats = vectorstore.similarity_search(question, k=2)
    return "\n---\n".join(doc.page_content for doc in resultats)


@tool
def generer_rapport(type_rapport: str):
    """Genere un rapport Word (docx) avec les donnees actuelles.
    type_rapport doit etre exactement l''un de : stock, ventes, factures."""
    if type_rapport == "stock":
        titre = "Rapport de stock"
        donnees = get_stock(limit=20)
    elif type_rapport == "ventes":
        titre = "Rapport des ventes"
        donnees = get_sales(limit=20)
    elif type_rapport == "factures":
        titre = "Rapport des factures"
        donnees = get_invoices(limit=20)
    else:
        return "Type de rapport non reconnu. Utilise stock, ventes ou factures."

    doc = Document()
    doc.add_heading(titre, level=1)
    for ligne in donnees:
        doc.add_paragraph(str(ligne))

    Path("rapports").mkdir(exist_ok=True)
    nom_fichier = f"{type_rapport}.docx"
    doc.save(f"rapports/{nom_fichier}")
    return f"Rapport genere avec succes. Telechargeable sur /rapports/{nom_fichier}"


llm = ChatGroq(
    model="openai/gpt-oss-120b",
    temperature=0,
    api_key=os.environ["GROQ_API_KEY"]
)

llm_with_tools = llm.bind_tools(
    [stock, sales, invoices, chercher_documents, generer_rapport]
)


def ask_agent(question):
    response = llm_with_tools.invoke(question)

    if not response.tool_calls:
        return response.content

    results = []
    for call in response.tool_calls:
        if call["name"] == "stock":
            result = stock.invoke(call["args"])
        elif call["name"] == "sales":
            result = sales.invoke(call["args"])
        elif call["name"] == "invoices":
            result = invoices.invoke(call["args"])
        elif call["name"] == "chercher_documents":
            result = chercher_documents.invoke(call["args"])
        elif call["name"] == "generer_rapport":
            result = generer_rapport.invoke(call["args"])
        results.append(str(result))

    final_prompt = f"""
Tu es un assistant de gestion d''entreprise.
Question :
{question}
Donnees recuperees :
{chr(10).join(results)}
Reponds en francais.
N''invente aucune information.
Utilise uniquement les donnees fournies.
Si un rapport a ete genere, mentionne le lien de telechargement tel quel.
"""
    final_response = llm.invoke(final_prompt)
    return final_response.content


if __name__ == "__main__":
    print("Assistant B Tel Bureautique -- tape quitter pour arreter\n")
    while True:
        question = input("Question : ")
        if question.strip().lower() in ("quitter", "exit", "quit"):
            print("A bientot !")
            break
        answer = ask_agent(question)
        print("\nReponse :")
        print(answer)
        print()
